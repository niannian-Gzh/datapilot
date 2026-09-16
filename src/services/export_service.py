import json
import pandas as pd
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from nl2sql import generate_filter_sql
from execute import execute_sql
from audit import log
from config import PROJECT_ROOT
import re


def _sanitize_filename(text: str, max_len: int = 30) -> str:
    """把用户输入变成合法的文件名片段。"""
    # 去掉 Windows 非法字符
    cleaned = re.sub(r'[\\/:*?"<>|]', "", text)
    # 去掉多余空格
    cleaned = re.sub(r"\s+", "_", cleaned.strip())
    # 截断
    if len(cleaned) > max_len:
        cleaned = cleaned[:max_len]
    return cleaned or "导出数据"

EXPORT_DIR = PROJECT_ROOT / "data" / "exports"

EXPORT_COLUMNS = {
    "project_name": "项目名称",
    "owner": "负责人",
    "status_raw": "状态",
    "issued_date": "项目下发时间",
    "certified_date": "项目下证时间",
    "category": "分类",
    "resubmit_name": "重提项目名",
    "reject_date": "（最新）打回时间",
    "resubmit_date": "（打回）提交时间",
    "reject_count": "打回次数",
}


def run(filter_question: str, ctx) -> str:
    try:
        sql = generate_filter_sql(filter_question, ctx.llm, ctx.schema, ctx.table_name)
        df = execute_sql(sql, ctx.db_path)
    except Exception as e:
        return json.dumps({"error": f"导出失败：{e}"}, ensure_ascii=False)

    if len(df) == 0:
        return json.dumps({
            "status": "empty",
            "message": "没有找到符合条件的记录，未生成文件。",
        }, ensure_ascii=False)

    # 只保留业务字段，并改成中文列名
    available = [c for c in EXPORT_COLUMNS if c in df.columns]
    df = df[available].rename(columns=EXPORT_COLUMNS)

    # 日期字段转字符串
    for col in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[col]):
            df[col] = df[col].dt.strftime("%Y-%m-%d").fillna("")

    # 写文件
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    # 文件名用筛选条件，清洗非法字符
    name_part = _sanitize_filename(filter_question)
    filename = f"{name_part}_{timestamp}.xlsx"
    file_path = EXPORT_DIR / filename

    df.to_excel(file_path, index=False)

    # 设置列宽，避免日期显示 ##### 
    wb = load_workbook(file_path)
    ws = wb.active
    for i, col_name in enumerate(df.columns, 1):
        values = [str(col_name)] + [str(v) for v in df[col_name]]
        max_len = max(len(v) for v in values) if values else 10
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)
    wb.save(file_path)

    log("export", {
        "trace_id": ctx.trace.id,
        "user_input": ctx.trace.data["question"],
        "file": str(file_path),
        "row_count": len(df),
    })

    return json.dumps({
        "status": "ok",
        "file_path": str(file_path.relative_to(PROJECT_ROOT)),
        "row_count": len(df),
        "answer": f"已导出 {len(df)} 条记录到 {file_path.relative_to(PROJECT_ROOT)}",
    }, ensure_ascii=False)